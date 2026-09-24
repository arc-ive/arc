"""Tests for upload text extraction (issue #299).

Company Brain stores text, so an upload has to become text before the
existing chunking, embedding and citation path can touch it. These cover
what is accepted, what is refused, and -- most importantly -- that
nothing is accepted silently.
"""

import io
import zipfile
import zlib

import pytest

from arc.services.document_extraction import (
    MAX_UPLOAD_BYTES,
    DocumentFormat,
    EmptyExtractionError,
    ExtractionError,
    FileTooLargeError,
    UnsupportedFormatError,
    detect_format,
    extract,
)

LEAVE_POLICY = (
    "Every permanent employee receives 26 days of paid annual leave "
    "per year, in addition to UK public holidays."
)


def _pdf(text: str) -> bytes:
    """A minimal PDF carrying a real text layer, built without a library.

    Deliberately hand-built rather than pulling in reportlab: a test
    dependency whose only job is to produce a fixture is a dependency the
    deployment then has to justify. pypdf reads this exactly as it reads
    a generated one.
    """
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF"
    ).encode()
    return bytes(out)


def _docx(paragraphs, table_rows=None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestFormatDetection:
    """Format comes from the bytes, never from the filename.

    A caller-supplied extension deciding which parser runs is how a
    parser gets handed something it was never meant to read.
    """

    def test_pdf_is_detected_by_signature(self):
        assert detect_format(b"%PDF-1.7\nrest") is DocumentFormat.PDF

    def test_a_pdf_named_txt_is_still_a_pdf(self):
        assert detect_format(b"%PDF-1.7\n", "notes.txt") is DocumentFormat.PDF

    def test_text_named_pdf_is_still_text(self):
        assert detect_format(b"just words", "report.pdf") is DocumentFormat.TEXT

    def test_markdown_is_distinguished_only_by_name(self):
        # Safe to trust the name here: both are read as text, so the
        # distinction carries no parsing risk.
        assert detect_format(b"# Title", "policy.md") is DocumentFormat.MARKDOWN
        assert detect_format(b"# Title", "policy.txt") is DocumentFormat.TEXT

    def test_a_plain_zip_is_not_a_word_document(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("hello.txt", "hi")
        with pytest.raises(UnsupportedFormatError, match="ZIP archive"):
            detect_format(buffer.getvalue(), "policy.docx")


class TestRefusals:
    """Nothing is accepted silently."""

    def test_images_are_refused_and_say_why_when_ocr_is_off(self):
        """OCR off is the default, and the refusal is unchanged by ADR-014."""
        png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
        with pytest.raises(UnsupportedFormatError) as exc:
            extract(png, "scan.png")
        # The message must name the real reason, not just "unsupported":
        # an operator needs to know OCR is the missing piece.
        assert "optical character recognition" in str(exc.value).lower()

    def test_binary_junk_is_refused(self):
        with pytest.raises(UnsupportedFormatError):
            extract(b"\x00\x01\x02\x03\xfe\xff" * 20, "mystery.bin")

    def test_oversized_files_are_refused_before_parsing(self):
        with pytest.raises(FileTooLargeError):
            extract(b"%PDF-1.7" + b"a" * (MAX_UPLOAD_BYTES + 1), "huge.pdf")

    def test_an_empty_file_is_refused(self):
        with pytest.raises(EmptyExtractionError):
            extract(b"", "empty.txt")

    def test_a_file_yielding_no_text_is_refused(self):
        """A stored document with no content retrieves as an empty
        citation and looks like an answer. Worse than no document."""
        with pytest.raises(EmptyExtractionError):
            extract(b"hi", "tiny.txt")

    def test_a_malformed_pdf_is_a_controlled_failure(self):
        """A parser exception must never escape as a 500."""
        with pytest.raises(ExtractionError):
            extract(b"%PDF-1.7\nnot actually a pdf at all" + b"\x00" * 200, "bad.pdf")


class TestExtraction:
    def test_plain_text_passes_through(self):
        result = extract(LEAVE_POLICY.encode(), "leave.txt")
        assert result.detected_format is DocumentFormat.TEXT
        assert "26 days" in result.text
        assert result.char_count == len(result.text)

    def test_markdown_keeps_its_content(self):
        source = f"# Leave Policy\n\n{LEAVE_POLICY}"
        result = extract(source.encode(), "leave.md")
        assert result.detected_format is DocumentFormat.MARKDOWN
        assert "26 days" in result.text

    def test_word_paragraphs_are_extracted(self):
        result = extract(_docx([LEAVE_POLICY]), "leave.docx")
        assert result.detected_format is DocumentFormat.DOCX
        assert "26 days" in result.text

    def test_word_tables_are_extracted(self):
        """Policy documents put entitlements in tables.

        Paragraph-only extraction silently drops them, so the fact
        someone will actually ask about is the fact that goes missing.
        """
        data = _docx(
            ["Notice required:"],
            table_rows=[["Length", "Notice"], ["3-10 days", "2 weeks"]],
        )
        result = extract(data, "notice.docx")
        assert "3-10 days" in result.text
        assert "2 weeks" in result.text

    def test_pdf_text_is_extracted(self):
        result = extract(_pdf(LEAVE_POLICY), "leave.pdf")
        assert result.detected_format is DocumentFormat.PDF
        assert "26 days" in result.text


# ---------------------------------------------------------------------------
# Optical character recognition (issue #299, ADR-014)
# ---------------------------------------------------------------------------


def _png(size=(400, 120), colour=255) -> bytes:
    """A real PNG, produced by the library pypdf itself decodes with."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("L", size, colour).save(buffer, format="PNG")
    return buffer.getvalue()


def _scanned_pdf(pages: int = 1) -> bytes:
    """A PDF with NO text layer whose every page is one embedded image.

    This is what a scanner produces, and it is the case OCR exists for.
    Built by hand for the same reason the text-layer fixture above is: a
    dependency whose only job is to produce a fixture is a dependency the
    deployment then has to justify.

    The images are page-shaped (1000x1400) because the extractor filters
    embedded images by DIMENSION -- a logo is small, a page is not -- so
    an icon-sized fixture would be skipped and prove nothing.
    """
    raster = zlib.compress(b"\xff" * (1000 * 1400))

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        None,  # pages node, filled once the kids are known
    ]
    kids = []
    for index in range(pages):
        page_obj = len(objects) + 1
        image_obj = page_obj + 1
        content_obj = page_obj + 2
        kids.append(page_obj)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 1000 1400] /Resources "
            b"<< /XObject << /Im0 "
            + str(image_obj).encode()
            + b" 0 R >> >> /Contents "
            + str(content_obj).encode()
            + b" 0 R >>"
        )
        objects.append(
            b"<< /Type /XObject /Subtype /Image /Width 1000 /Height 1400 "
            b"/ColorSpace /DeviceGray /BitsPerComponent 8 /Filter /FlateDecode "
            b"/Length " + str(len(raster)).encode() + b" >>\nstream\n" + raster + b"\nendstream"
        )
        content = b"q 1000 0 0 1400 0 0 cm /Im0 Do Q"
        objects.append(
            b"<< /Length "
            + str(len(content)).encode()
            + b" >>\nstream\n"
            + content
            + b"\nendstream"
        )

    objects[1] = (
        b"<< /Type /Pages /Kids ["
        + b" ".join(f"{kid} 0 R".encode() for kid in kids)
        + b"] /Count "
        + str(pages).encode()
        + b" >>"
    )

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def _pdf_with_text_and_image(text: str) -> bytes:
    """A PDF that has BOTH a real text layer and a page-sized image.

    This exists so the "never re-read a text layer with OCR" test can
    actually fail. With a text-only PDF it cannot: there is no image to
    read instead, so the assertion holds whether or not the preference is
    implemented. A scanned page with a searchable-text overlay is also a
    real shape -- it is what an OCR'd archive looks like.
    """
    raster = zlib.compress(b"\xff" * (1000 * 1400))
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    stream += b"\nq 1000 0 0 1400 0 0 cm /Im0 Do Q"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 1000 1400] /Resources "
        b"<< /Font << /F1 4 0 R >> /XObject << /Im0 6 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /XObject /Subtype /Image /Width 1000 /Height 1400 "
        b"/ColorSpace /DeviceGray /BitsPerComponent 8 /Filter /FlateDecode "
        b"/Length " + str(len(raster)).encode() + b" >>\nstream\n" + raster + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


class _FakeOcr:
    """A recogniser that returns fixed text, and records what it saw.

    Standing in for Tesseract, not for the pipeline: everything between
    the uploaded bytes and this call is the real code, including format
    detection, the text-layer-first rule and the PDF image extraction.
    """

    name = "fake"

    def __init__(self, text=LEAVE_POLICY, fail=False):
        self._text = text
        self._fail = fail
        self.seen = []

    def extract_text(self, image: bytes) -> str:
        self.seen.append(image)
        if self._fail:
            from arc.services.ocr import OcrError

            raise OcrError("engine unavailable")
        return self._text


class TestOcrOffByDefault:
    """ADR-014: a deployment that does not opt in is unchanged."""

    def test_an_image_is_still_refused(self):
        with pytest.raises(UnsupportedFormatError, match="(?i)optical character recognition"):
            extract(_png(), "scan.png")

    def test_a_scanned_pdf_is_still_refused_and_names_ocr(self):
        with pytest.raises(EmptyExtractionError) as exc:
            extract(_scanned_pdf(), "scan.pdf")
        assert "optical character recognition" in str(exc.value).lower()

    def test_a_text_layer_pdf_is_unaffected(self):
        result = extract(_pdf(LEAVE_POLICY), "policy.pdf")
        assert result.ocr_used is False
        assert result.ocr_provider is None


class TestOcrReadsImages:
    def test_an_image_becomes_a_document(self):
        ocr = _FakeOcr()

        result = extract(_png(), "scan.png", ocr=ocr)

        assert result.detected_format is DocumentFormat.IMAGE
        assert LEAVE_POLICY in result.text
        assert result.ocr_used is True
        assert result.ocr_provider == "fake"
        assert len(ocr.seen) == 1

    def test_an_image_that_reads_as_nothing_is_still_refused(self):
        """Storing an empty document is the thing this must never do."""
        with pytest.raises(EmptyExtractionError) as exc:
            extract(_png(), "blank.png", ocr=_FakeOcr(text="   "))
        assert "found no text" in str(exc.value).lower()

    def test_an_engine_failure_does_not_become_a_stored_empty_document(self):
        with pytest.raises(EmptyExtractionError):
            extract(_png(), "scan.png", ocr=_FakeOcr(fail=True))


class TestOcrReadsScannedPdfs:
    def test_a_scanned_pdf_is_read_through_its_page_image(self):
        ocr = _FakeOcr()
        pdf = _scanned_pdf()

        result = extract(pdf, "scan.pdf", ocr=ocr)

        assert result.detected_format is DocumentFormat.PDF
        assert LEAVE_POLICY in result.text
        assert result.ocr_used is True
        assert len(ocr.seen) == 1

    def test_a_pdf_with_a_text_layer_is_never_re_read_by_ocr(self):
        """The text layer is exact; OCR is a guess. Never replace one with
        the other, and never pay for OCR on a document that has text.

        The fixture carries a text layer AND a page-sized image, so
        removing the preference would visibly substitute the recognised
        text for the real one. A text-only PDF could not detect that.
        """
        ocr = _FakeOcr(text="THIS SHOULD NOT APPEAR")

        result = extract(_pdf_with_text_and_image(LEAVE_POLICY), "overlay.pdf", ocr=ocr)

        assert "26 days" in result.text
        assert "THIS SHOULD NOT APPEAR" not in result.text
        assert result.ocr_used is False
        assert ocr.seen == []

    def test_the_number_of_images_read_is_bounded(self):
        """A long scan must not occupy a worker without limit."""
        ocr = _FakeOcr()

        result = extract(_scanned_pdf(pages=5), "long-scan.pdf", ocr=ocr, max_ocr_images=2)

        assert len(ocr.seen) == 2
        assert result.ocr_used is True

    def test_every_page_of_a_short_scan_is_read(self):
        """The bound is a ceiling, not a default truncation."""
        ocr = _FakeOcr()

        extract(_scanned_pdf(pages=3), "scan.pdf", ocr=ocr, max_ocr_images=20)

        assert len(ocr.seen) == 3

    def test_a_pdf_with_no_text_and_no_images_is_refused(self):
        """A vector-only PDF has nothing for OCR to read, and says so."""
        ocr = _FakeOcr()
        with pytest.raises(EmptyExtractionError):
            extract(_pdf(""), "vector.pdf", ocr=ocr)
        assert ocr.seen == []


class TestOcrSettings:
    def test_ocr_is_off_unless_configured(self, monkeypatch):
        from arc.services.ocr import build_ocr_provider

        monkeypatch.delenv("OCR_PROVIDER", raising=False)
        assert build_ocr_provider() is None

    @pytest.mark.parametrize("value", ["none", "None", " off ", "disabled", ""])
    def test_explicit_off_values_are_honoured(self, monkeypatch, value):
        from arc.services.ocr import build_ocr_provider

        monkeypatch.setenv("OCR_PROVIDER", value)
        assert build_ocr_provider() is None

    def test_an_unknown_provider_is_refused_rather_than_ignored(self, monkeypatch):
        """Silently falling back to 'no OCR' would hide an operator typo."""
        from arc.services.ocr import OcrUnavailableError, build_ocr_provider

        monkeypatch.setenv("OCR_PROVIDER", "magic-cloud-ocr")
        with pytest.raises(OcrUnavailableError, match="not a supported provider"):
            build_ocr_provider()

    def test_tesseract_without_the_binary_fails_loudly_at_build_time(self, monkeypatch):
        """An operator who asked for OCR learns at startup, not at upload."""
        from arc.services.ocr import OcrUnavailableError, build_ocr_provider

        monkeypatch.setenv("OCR_PROVIDER", "tesseract")
        try:
            import pytesseract

            has_binary = True
            try:
                pytesseract.get_tesseract_version()
            except Exception:
                has_binary = False
        except ImportError:
            has_binary = False

        if has_binary:
            pytest.skip("tesseract is installed here; the failure path cannot be exercised")
        with pytest.raises(OcrUnavailableError):
            build_ocr_provider()

    @pytest.mark.parametrize(
        "raw,expected",
        [(None, 20), ("", 20), ("5", 5), ("0", 20), ("-3", 20), ("not-a-number", 20)],
    )
    def test_bounds_fall_back_rather_than_crash(self, monkeypatch, raw, expected):
        from arc.services.ocr import get_ocr_settings

        if raw is None:
            monkeypatch.delenv("OCR_MAX_PAGES", raising=False)
        else:
            monkeypatch.setenv("OCR_MAX_PAGES", raw)
        assert get_ocr_settings()["max_pages"] == expected
